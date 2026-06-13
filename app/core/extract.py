"""Extração de texto de contratos em PDF e DOCX."""
from __future__ import annotations

from pathlib import Path

from PyPDF2 import PdfReader
import docx

SUPORTADOS = {".pdf", ".docx"}


def extract_pdf(path: Path) -> str:
    """Extrai texto de um PDF, página a página."""
    reader = PdfReader(str(path))
    partes: list[str] = []
    for pagina in reader.pages:
        texto = pagina.extract_text() or ""
        if texto.strip():
            partes.append(texto)
    return "\n\n".join(partes).strip()


def extract_docx(path: Path) -> str:
    """Extrai texto de um DOCX, incluindo parágrafos e células de tabelas."""
    documento = docx.Document(str(path))
    partes: list[str] = [p.text for p in documento.paragraphs if p.text.strip()]
    for tabela in documento.tables:
        for linha in tabela.rows:
            celulas = [c.text.strip() for c in linha.cells if c.text.strip()]
            if celulas:
                partes.append(" | ".join(celulas))
    return "\n".join(partes).strip()


def extract_text(path: Path) -> str:
    """Despacha a extração consoante a extensão do ficheiro."""
    sufixo = path.suffix.lower()
    if sufixo == ".pdf":
        return extract_pdf(path)
    if sufixo == ".docx":
        return extract_docx(path)
    # TXT/MD usados nas referências.
    if sufixo in {".txt", ".md"}:
        return path.read_text(encoding="utf-8", errors="replace").strip()
    raise ValueError(f"Formato não suportado: {path.name} ({sufixo})")
