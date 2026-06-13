"""Geração do relatório DOCX a partir da análise."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

from app.core.models import AnaliseContrato

_CORES_SEVERIDADE = {
    "alta": RGBColor(0xC0, 0x00, 0x00),
    "média": RGBColor(0xB8, 0x6A, 0x00),
    "baixa": RGBColor(0x1F, 0x6F, 0x1F),
}


def _add_bullets(doc: Document, itens: list[str]) -> None:
    if not itens:
        doc.add_paragraph("—")
        return
    for item in itens:
        doc.add_paragraph(str(item), style="List Bullet")


def gerar_docx(
    analise: AnaliseContrato,
    destino: Path,
    nome_contrato: str,
    tipo_nome: str = "",
) -> None:
    """Constrói e grava o relatório DOCX em `destino`."""
    doc = Document()

    titulo = doc.add_heading("Relatório de Análise de Contrato", level=0)
    titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if tipo_nome:
        tp = doc.add_paragraph(f"Tipo: {tipo_nome}")
        tp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        tp.runs[0].bold = True
    sub = doc.add_paragraph(nome_contrato)
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.runs[0].italic = True

    # --- Resumo ---
    r = analise.resumo
    doc.add_heading("1. Resumo", level=1)
    if r.titulo:
        p = doc.add_paragraph()
        p.add_run("Título/Tipo: ").bold = True
        p.add_run(r.titulo)
    if r.cliente:
        p = doc.add_paragraph()
        p.add_run("Cliente: ").bold = True
        p.add_run(r.cliente)
    if r.objeto:
        p = doc.add_paragraph()
        p.add_run("Objeto: ").bold = True
        p.add_run(r.objeto)

    doc.add_heading("Partes", level=2)
    _add_bullets(doc, r.partes)
    doc.add_heading("Valores", level=2)
    _add_bullets(doc, r.valores)
    doc.add_heading("Prazos", level=2)
    _add_bullets(doc, r.prazos)
    if r.sintese:
        doc.add_heading("Síntese", level=2)
        doc.add_paragraph(r.sintese)

    # --- Riscos ---
    doc.add_heading("2. Riscos e Alertas", level=1)
    if not analise.riscos:
        doc.add_paragraph("Não foram identificados riscos relevantes.")
    else:
        # Ordena por severidade (alta -> baixa).
        ordem = {"alta": 0, "média": 1, "baixa": 2}
        for risco in sorted(analise.riscos, key=lambda x: ordem.get(x.severidade, 3)):
            p = doc.add_paragraph(style="List Bullet")
            run_sev = p.add_run(f"[{risco.severidade.upper()}] ")
            run_sev.bold = True
            run_sev.font.color.rgb = _CORES_SEVERIDADE.get(
                risco.severidade, RGBColor(0, 0, 0)
            )
            p.add_run(risco.descricao)
            if risco.clausula:
                det = doc.add_paragraph()
                det.paragraph_format.left_indent = Pt(24)
                det.add_run("Cláusula: ").italic = True
                det.add_run(risco.clausula)
            if risco.recomendacao:
                det = doc.add_paragraph()
                det.paragraph_format.left_indent = Pt(24)
                det.add_run("Recomendação: ").italic = True
                det.add_run(risco.recomendacao)

    destino.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(destino))
